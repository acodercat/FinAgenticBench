import csv
import importlib
import json
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
import pytest
from openpyxl import load_workbook

from config import PROJECT_ROOT
from core.data import load_expansion_table, runtime_variables
from core.dataset_layers import RAW_DIR


def _expected_outputs(module):
    return {
        variable.name: value
        for variable, value in zip(module.variables, module.ground_truth())
    }


def _claims(spec):
    """Every claim in the case, in turn order: a case may spread them over turns."""
    return [claim for turn in spec["conversations"][0]["turns"] for claim in turn["claims"]]


def test_phase9_build_report_and_governed_runtime_tables():
    report = json.loads(
        (PROJECT_ROOT / "metadata/build_reports/phase9_cross_source_table_build.json").read_text()
    )
    assert report["eia_steo_wti_vintages"] == {
        "rows": 576,
        "publication_vintages": 24,
        "target_months": 24,
        "output": "eia_steo_vintages/wti_monthly_vintages.csv.gz",
    }
    assert report["eia_pet_wti_monthly_actuals"]["rows"] == 480
    assert report["eia_pet_wti_monthly_actuals"]["series_id"] == "PET.RWTC.M"
    assert report["eia_pet_wti_daily_actuals"]["rows"] == 10091
    assert report["eia_pet_wti_daily_actuals"]["series_id"] == "PET.RWTC.D"
    assert report["fema_declaration_areas"]["rows"] == 3567
    assert report["fema_declaration_areas"]["disaster_numbers"] == 301

    vintages = load_expansion_table("eia_steo_vintages")
    actuals = load_expansion_table("eia_bulk_wti")
    daily_actuals = load_expansion_table("eia_bulk_wti_daily")
    declarations = load_expansion_table("fema_disasters")
    assert not vintages.duplicated(["publication_date", "target_month"]).any()
    assert actuals.observation_month.is_unique
    assert actuals.observation_month.max() == "2025-12"
    assert daily_actuals.observation_date.is_unique
    assert daily_actuals.observation_date.max() == "2025-12-31"
    assert declarations.declaration_area_id.is_unique
    assert declarations.county_fips.str.fullmatch(r"\d{5}").all()


def test_phase9_runtime_exposes_only_the_requested_local_sources():
    variables = runtime_variables(
        {"eia_steo_vintages", "eia_bulk_wti", "fema_disasters", "fdic_sod"}
    )
    # Exposure is what the declaration guarantees; the sequence is registry order.
    assert sorted(variable.name for variable in variables) == sorted([
        "fdic_sod_branches_df",
        "eia_steo_wti_vintages_df",
        "eia_wti_monthly_spot_prices_df",
        "eia_wti_daily_spot_prices_df",
        "fema_declaration_areas_df",
    ])
    descriptions = {variable.name: variable.description for variable in variables}
    assert "later-vintage history must not replace an earlier forecast" in descriptions[
        "eia_steo_wti_vintages_df"
    ]
    assert "current revised bulk snapshot" in descriptions[
        "eia_wti_monthly_spot_prices_df"
    ]
    assert "non-observation dates are not synthesized" in descriptions[
        "eia_wti_daily_spot_prices_df"
    ]
    assert "A county can appear in multiple disaster declarations" in descriptions[
        "fema_declaration_areas_df"
    ]


def test_wti_case_anchor_values_and_validator():
    module = importlib.import_module(
        "evals.temporal_path.wti_forecast_vintage_error_deterioration"
    )
    truth = dict(zip((variable.name for variable in module.variables), module.ground_truth()))
    assert truth["common_future_target_month_count"] == 8
    assert round(truth["january_vintage_mean_absolute_error_usd_per_barrel"], 3) == 3.951
    assert round(truth["april_vintage_mean_absolute_error_usd_per_barrel"], 3) == 10.929
    assert truth["april_lower_absolute_error_month_count"] == 0
    assert truth["largest_absolute_error_deterioration_month"] == "2024-08"
    assert (
        truth["selected_month_january_forecast_usd_per_barrel"],
        truth["selected_month_april_forecast_usd_per_barrel"],
        truth["selected_month_actual_usd_per_barrel"],
    ) == (77.50, 87.50, 76.68)
    assert truth["selected_month_daily_observation_count"] == 22
    assert (
        truth["selected_month_daily_low_usd_per_barrel"],
        truth["selected_month_daily_high_usd_per_barrel"],
    ) == (72.76, 81.45)
    expected = _expected_outputs(module)
    assert module.validate(expected).success
    assert not module.validate(
        {**expected, "selected_month_actual_usd_per_barrel": 76.86}
    ).success


def test_fema_sod_case_anchor_values_and_validator():
    module = importlib.import_module(
        "evals.risk_screening.fema_sod_disaster_deposit_exposure"
    )
    truth = module.ground_truth()
    assert truth[:4] == (19, 359, "Bank of America, National Association", 3510)
    assert tuple(round(value, 6) for value in truth[4:6]) == (
        124.117307,
        1975.044000,
    )
    assert round(truth[6], 4) == 6.2843
    assert truth[7:9] == ("City National Bank", 17281)
    assert tuple(round(value, 6) for value in truth[9:11]) == (
        50.241248,
        77.761501,
    )
    assert round(truth[11], 4) == 64.6094
    assert round(truth[12], 4) == 46.9697
    expected = _expected_outputs(module)
    assert module.validate(
        {
            **expected,
            "largest_exposed_deposit_amount_bank_name": " bank of america, national association ",
            "largest_exposed_deposit_share_bank_name": "CITY NATIONAL BANK",
        }
    ).success
    assert not module.validate(
        {**expected, "distinct_qualifying_county_count": 478}
    ).success


def test_phase9_source_removal_and_claim_contracts_are_explicit():
    registry = json.loads((PROJECT_ROOT / "benchmarks.json").read_text())
    paths = registry["cases"]
    assert paths["wti_forecast_vintage_error_deterioration"] == (
        "evals/temporal_path/wti_forecast_vintage_error_deterioration.json"
    )
    assert paths["fema_sod_disaster_deposit_exposure"] == (
        "evals/risk_screening/fema_sod_disaster_deposit_exposure.json"
    )
    wti = json.loads(
        (
            PROJECT_ROOT
            / "evals/temporal_path/wti_forecast_vintage_error_deterioration.json"
        ).read_text()
    )
    fema = json.loads(
        (
            PROJECT_ROOT
            / "evals/risk_screening/fema_sod_disaster_deposit_exposure.json"
        ).read_text()
    )
    assert wti["data_sources"] == ["eia_steo_vintages", "eia_bulk_wti"]
    assert fema["data_sources"] == ["fema_disasters", "fdic_sod"]
    wti_modes = {claim["id"]: claim["mode"] for claim in _claims(wti)}
    assert wti_modes["later_vintage_accuracy_result"] == "required"
    assert wti_modes["forecast_actual_vintage_boundary"] == "required"
    assert wti_modes["monthly_daily_measure_boundary"] == "required"
    assert wti_modes["revision_causality_boundary"] == "conditional"
    fema_modes = {claim["id"]: claim["mode"] for claim in _claims(fema)}
    assert fema_modes["amount_share_leader_divergence"] == "required"
    assert fema_modes["sod_geographic_measure_boundary"] == "required"
    assert fema_modes["fema_declaration_measure_boundary"] == "required"
    assert fema_modes["retrospective_vintage_boundary"] == "conditional"


def _workbook_wti_values(path: Path) -> dict[str, float]:
    worksheet = load_workbook(path, read_only=True, data_only=True)["2tab"]
    months = {
        name: number
        for number, name in enumerate(
            ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"],
            start=1,
        )
    }
    year = None
    columns = {}
    for index, (year_value, month_value) in enumerate(
        zip(worksheet[3], worksheet[4]), start=1
    ):
        if isinstance(year_value.value, (int, float)):
            year = int(year_value.value)
        month = months.get(str(month_value.value)[:3]) if month_value.value else None
        if year == 2024 and month in range(5, 13):
            columns[f"2024-{month:02d}"] = index
    row = next(
        values
        for values in worksheet.iter_rows(min_row=5, values_only=True)
        if values[0] == "WTIPUUS"
    )
    return {month: float(row[index - 1]) for month, index in columns.items()}


@pytest.mark.raw_data
def test_independent_raw_recomputation_matches_wti_case():
    january = _workbook_wti_values(
        RAW_DIR / "eia_steo_vintages/jan24_base.xlsx"
    )
    april = _workbook_wti_values(
        RAW_DIR / "eia_steo_vintages/apr24_base.xlsx"
    )
    with ZipFile(RAW_DIR / "eia_bulk/PET.zip") as archive:
        with archive.open("PET.txt") as handle:
            sources = {
                source["series_id"]: source
                for raw in handle
                if (source := json.loads(raw)).get("series_id")
                in {"PET.RWTC.M", "PET.RWTC.D"}
            }
    source = sources["PET.RWTC.M"]
    actual = {
        f"{period[:4]}-{period[4:]}": float(value)
        for period, value in source["data"]
        if "202405" <= period <= "202412"
    }
    rows = []
    for month in sorted(actual):
        january_error = january[month] - actual[month]
        april_error = april[month] - actual[month]
        rows.append(
            (
                month,
                january[month],
                april[month],
                actual[month],
                january_error,
                april_error,
                abs(april_error) - abs(january_error),
            )
        )
    leader = sorted(rows, key=lambda row: (-row[6], row[0]))[0]
    selected_daily = [
        float(value)
        for period, value in sources["PET.RWTC.D"]["data"]
        if period.startswith(leader[0].replace("-", "")) and value is not None
    ]
    independent = (
        len(rows),
        sum(abs(row[4]) for row in rows) / len(rows),
        sum(abs(row[5]) for row in rows) / len(rows),
        sum(abs(row[5]) < abs(row[4]) for row in rows),
        *leader[:4],
        len(selected_daily),
        min(selected_daily),
        max(selected_daily),
    )
    module = importlib.import_module(
        "evals.temporal_path.wti_forecast_vintage_error_deterioration"
    )
    truth = module.ground_truth()
    assert independent[:5] == truth[:5]
    assert all(abs(a - b) < 1e-12 for a, b in zip(independent[5:], truth[5:]))


@pytest.mark.raw_data
def test_independent_raw_recomputation_matches_fema_sod_case():
    declarations = pd.read_parquet(
        RAW_DIR / "fema_disasters/DisasterDeclarationsSummaries.parquet"
    )
    declarations["declarationDate"] = pd.to_datetime(declarations.declarationDate)
    declarations["incidentBeginDate"] = pd.to_datetime(declarations.incidentBeginDate)
    declarations["county_fips"] = (
        declarations.fipsStateCode.astype(str).str.zfill(2)
        + declarations.fipsCountyCode.astype(str).str.zfill(3)
    )
    qualifying = declarations.loc[
        declarations.declarationType.eq("DR")
        & declarations.declarationDate.le("2025-06-30")
        & declarations.incidentBeginDate.between("2025-01-01", "2025-06-30")
        & declarations.fipsCountyCode.ne("000")
    ]
    county_fips = set(qualifying.county_fips)
    banks: dict[tuple[int, str], dict[str, float]] = {}
    with (
        RAW_DIR / "fdic_sod/sod_2025_all.csv"
    ).open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            key = (int(row["CERT"]), row["NAMEFULL"].strip())
            bank = banks.setdefault(
                key, {"rows": 0, "deposits": 0, "exposed_rows": 0, "exposed": 0}
            )
            deposit = int(row["DEPSUMBR"])
            exposed = row["STCNTYBR"].zfill(5) in county_fips
            bank["rows"] += 1
            bank["deposits"] += deposit
            bank["exposed_rows"] += int(exposed)
            bank["exposed"] += deposit if exposed else 0
    top = sorted(
        banks.items(), key=lambda item: (-item[1]["deposits"], item[0][1])
    )[:50]
    amount = sorted(top, key=lambda item: (-item[1]["exposed"], item[0][1]))[0]
    share = sorted(
        top,
        key=lambda item: (
            -(item[1]["exposed"] / item[1]["deposits"]),
            item[0][1],
        ),
    )[0]
    independent = (
        qualifying.disasterNumber.nunique(),
        qualifying.county_fips.nunique(),
        amount[0][1],
        amount[0][0],
        amount[1]["exposed"] / 1e6,
        amount[1]["deposits"] / 1e6,
        amount[1]["exposed"] / amount[1]["deposits"] * 100,
        share[0][1],
        share[0][0],
        share[1]["exposed"] / 1e6,
        share[1]["deposits"] / 1e6,
        share[1]["exposed"] / share[1]["deposits"] * 100,
        share[1]["exposed_rows"] / share[1]["rows"] * 100,
    )
    module = importlib.import_module(
        "evals.risk_screening.fema_sod_disaster_deposit_exposure"
    )
    truth = module.ground_truth()
    assert independent[:4] == truth[:4]
    assert independent[7:9] == truth[7:9]
    numeric_pairs = list(zip(independent[4:7], truth[4:7])) + list(
        zip(independent[9:], truth[9:])
    )
    assert all(abs(a - b) < 1e-12 for a, b in numeric_pairs)
